"""
=============================================================================
FICHIER : backend/audit_data.py
ROLE    : Audit des fichiers sources avant import en base de donnees

Ce script analyse les fichiers Excel et le Shapefile pour comprendre leur
structure, detecter les problemes et preparer les commandes d'import.

Questions auxquelles ce script repond :
  - Quelles colonnes existent dans chaque fichier Excel ?
  - Combien de paroisses ont des coordonnees GPS ? Combien n'en ont pas ?
  - Quels sont les noms exacts des regions dans le shapefile ?
  - Quels types d'oeuvres existent dans les donnees ?
  - Quels grades d'ouvriers sont presents ?

COMMANDE D'EXECUTION (depuis la racine du projet) :
  docker compose run --rm -v "d:/Academique/GEOEEC/data:/data" backend python audit_data.py
=============================================================================
"""
import os
import sys
from pathlib import Path

# Initialiser Django pour utiliser son wrapper GDAL (lecture du shapefile)
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "eec_core.settings.dev")

import django
django.setup()

try:
    import openpyxl
except ImportError:
    print("ERREUR: openpyxl non installe")
    sys.exit(1)

SEPARATEUR = "=" * 70


def auditer_excel(chemin: str, nom: str):
    """Analyse complète d'un fichier Excel."""
    print(f"\n{SEPARATEUR}")
    print(f"FICHIER : {nom}")
    print(SEPARATEUR)

    try:
        wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    except Exception as e:
        print(f"  ERREUR ouverture : {e}")
        return

    print(f"  Feuilles dans le fichier : {wb.sheetnames}")

    for nom_feuille in wb.sheetnames:
        ws = wb[nom_feuille]
        print(f"\n  --- Feuille : '{nom_feuille}' ---")

        # Lire toutes les lignes en mémoire
        lignes = list(ws.iter_rows(values_only=True))
        if not lignes:
            print("  Feuille vide.")
            continue

        entetes = lignes[0]
        donnees = lignes[1:]

        print(f"  Nombre de colonnes : {len(entetes)}")
        print(f"  Nombre de lignes de données : {len(donnees)}")
        print(f"\n  COLONNES DÉTECTÉES :")

        for i, col in enumerate(entetes):
            if col is None:
                continue

            # Compter valeurs non-nulles, nulles, et types
            valeurs = [row[i] for row in donnees if i < len(row)]
            non_nulles = [v for v in valeurs if v is not None and str(v).strip() != ""]
            nulles = len(valeurs) - len(non_nulles)

            # Détecter le type dominant
            types = set(type(v).__name__ for v in non_nulles[:50])
            type_str = ", ".join(types) if types else "vide"

            print(f"\n  [{i+1}] Colonne : '{col}'")
            print(f"       Remplies : {len(non_nulles)} / {len(valeurs)}  |  Vides : {nulles}")
            print(f"       Type de données : {type_str}")

            # Afficher les valeurs uniques si peu nombreuses (champ catégoriel)
            uniques = list(set(str(v) for v in non_nulles if v is not None))
            if 1 < len(uniques) <= 30:
                uniques_trie = sorted(uniques)
                print(f"       Valeurs uniques ({len(uniques)}) : {uniques_trie}")
            elif len(uniques) == 1:
                print(f"       Valeur unique : '{uniques[0]}'")
            elif len(uniques) > 30:
                # Afficher quelques exemples
                exemples = sorted(uniques)[:5]
                print(f"       Trop de valeurs uniques ({len(uniques)}) — exemples : {exemples}")

        # Chercher colonnes GPS spécifiquement
        print(f"\n  ANALYSE GPS :")
        cols_lower = [str(c).lower() if c else "" for c in entetes]
        for mot in ["coord", "lat", "lon", "gps", "x", "y", "longitude", "latitude"]:
            for i, c in enumerate(cols_lower):
                if mot in c and entetes[i] not in [None, ""]:
                    vals = [row[i] for row in donnees if i < len(row) and row[i] is not None and str(row[i]).strip() != ""]
                    print(f"    Colonne GPS trouvée : '{entetes[i]}' → {len(vals)} valeurs non-vides")
                    if vals:
                        exemples = vals[:3]
                        print(f"      Exemples : {exemples}")

    wb.close()


def auditer_shapefile(chemin_shp: str):
    """
    Analyse le Shapefile des regions synodales via le wrapper GDAL de Django.

    Affiche :
      - Le systeme de coordonnees (CRS)
      - Le type de geometrie (MultiPolygon)
      - Les champs disponibles dans la table attributaire
      - Toutes les valeurs de chaque champ (pour voir les noms de regions)
    """
    print(f"\n{SEPARATEUR}")
    print(f"SHAPEFILE : {chemin_shp}")
    print(SEPARATEUR)

    # On utilise le DataSource de Django (wrapper autour de GDAL)
    # Pas besoin d'importer osgeo directement — Django l'expose proprement
    from django.contrib.gis.gdal import DataSource

    try:
        ds = DataSource(chemin_shp)
    except Exception as e:
        print(f"  ERREUR ouverture shapefile : {e}")
        return

    # Un shapefile = 1 seule couche (layer index 0)
    layer = ds[0]

    print(f"  Nombre d'entites (polygones) : {len(layer)}")
    print(f"  Type de geometrie            : {layer.geom_type}")
    print(f"  Systeme de coordonnees       : {layer.srs}")
    print(f"  Champs disponibles           : {layer.fields}")

    # Afficher toutes les valeurs de chaque champ pour reperer les noms de regions
    print(f"\n  VALEURS DE CHAQUE CHAMP :")
    for champ in layer.fields:
        valeurs = sorted(set(str(feat[champ].value) for feat in layer if feat[champ].value))
        print(f"\n  Champ '{champ}' — {len(valeurs)} valeurs uniques :")
        for v in valeurs:
            print(f"    -> '{v}'")

    print(f"\n  Shapefile analyse avec succes.")


if __name__ == "__main__":
    print("\n" + "=" * 70)
    print("  AUDIT DES DONNÉES SOURCES — Projet EEC Géolocalisation")
    print("=" * 70)

    DATA_DIR = Path("/data")

    # --- Excel paroisses ---
    auditer_excel(
        str(DATA_DIR / "Recap_paroisses_Projet_de_géolocalisation_26mai.xlsx"),
        "Paroisses (Recap_paroisses...)"
    )

    # --- Excel ouvriers ---
    auditer_excel(
        str(DATA_DIR / "OUVRIERS.xlsx"),
        "Ouvriers (OUVRIERS.xlsx)"
    )

    # --- Excel oeuvres ---
    auditer_excel(
        str(DATA_DIR / "Recap_oeuvres_EEC_2025.xlsx"),
        "Oeuvres (Recap_oeuvres_EEC_2025.xlsx)"
    )

    # --- Shapefile régions synodales ---
    auditer_shapefile("/data/gis/Region_synodale_ok2.shp")

    print(f"\n{SEPARATEUR}")
    print("  AUDIT TERMINÉ")
    print(SEPARATEUR)
