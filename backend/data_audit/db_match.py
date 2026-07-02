# -*- coding: utf-8 -*-
"""Audit phase 3 — état de la BD + croisement avec les nouveaux fichiers."""
import json
import os
import re
import sys
import unicodedata

sys.stdout.reconfigure(encoding="utf-8")
import django  # noqa: E402

os.environ.setdefault("DJANGO_SETTINGS_MODULE", os.environ.get("DJANGO_SETTINGS_MODULE", "eec_core.settings.dev"))
django.setup()

from apps.geo.models import Paroisse, RegionSynodale, District          # noqa: E402
from apps.oeuvres.models import Oeuvre                                   # noqa: E402
from apps.ouvriers.models import Ouvrier                                 # noqa: E402
from apps.accounts.models import StatistiqueAnnuelle as StatistiqueParoisse  # noqa: E402

D = "/app/data_audit/"


def norm(s: str) -> str:
    """Normalisation agressive pour la correspondance de noms."""
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode()
    s = s.lower()
    s = re.sub(r"\bparoisse( de| du| d')?\b", " ", s)
    s = re.sub(r"\beec\b", " ", s)
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return " ".join(s.split())


# ═══ 1. État de la BD ═══
print("═" * 70)
print("ÉTAT ACTUEL DE LA BASE DE DONNÉES")
print("═" * 70)
print(f"Régions synodales : {RegionSynodale.objects.count()}")
print(f"Districts         : {District.objects.count()}")
print(f"Paroisses         : {Paroisse.objects.count()}")
print(f"Œuvres            : {Oeuvre.objects.count()}")
print(f"Ouvriers          : {Ouvrier.objects.count()}")
print(f"Statistiques      : {StatistiqueParoisse.objects.count()} lignes "
      f"couvrant {StatistiqueParoisse.objects.values('paroisse').distinct().count()} paroisses")

print("\nChamps du modèle Paroisse :")
print("  " + ", ".join(f.name for f in Paroisse._meta.get_fields() if not f.is_relation or f.many_to_one))
print("Champs du modèle StatistiqueParoisse :")
print("  " + ", ".join(f.name for f in StatistiqueParoisse._meta.get_fields() if not f.is_relation or f.many_to_one))

db_par = list(Paroisse.objects.select_related("district__region").values(
    "id", "nom", "district__nom", "district__region__nom"))
db_index = {}
for p in db_par:
    db_index.setdefault(norm(p["nom"]), []).append(p)
dup_norm = {k: v for k, v in db_index.items() if len(v) > 1}
print(f"\nNoms de paroisse BD normalisés en doublon : {len(dup_norm)}")
for k, v in list(dup_norm.items())[:6]:
    print(f"   « {k} » ×{len(v)} : " + " | ".join(f"{x['nom']} ({x['district__region__nom']})" for x in v))

# ═══ 2. Croisement CATÉGORIES (docx 1 + 2) ═══
cat1 = json.load(open(D + "out_cat1.json", encoding="utf-8"))
cat2 = json.load(open(D + "out_cat2.json", encoding="utf-8"))
# docx1 : tables = C1, C2, C3, C4 ; paragraphes = A++/A1/A2/B1/B2 (extraits ici en dur)
AB = {
    "A++": ["Cinquantenaire", "Messa - Mokolo", "Biyem-Assi", "Nlongkak"],
    "A1": ["Bonamunduru", "Briqueterie 1", "Briqueterie 2", "Essos", "Mont des Oliviers",
           "New-Bell Aviation", "New-Deido", "Njimbam- Foumbot", "Paroisse de Bonabéri",
           "Paroisse de Bonatéki", "Paroisse de Centenaire (Act Fr)", "Paroisse Company",
           "Paroisse de Melen", "Paroisse de Soboum", "Paroisse Grand Temple", "Rue Manguiers"],
    "A2": ["Bafoussam Plateau", "Centre Martin Luther", "Maképè Tonnere", "Tamdja"],
    "B1": ["Bafang II", "Bamendzi A", "Bodomo", "Cité – Palmiers", "Koupgouo", "Messassi",
           "Ngousso Essesselakok", "Nkoldongo II", "Nkoldongo I", "Nkozoa", "Nsam Efoulan",
           "Nzintia", "Oyom Abang", "Tchitchap", "Tialong", "Toket"],
    "B2": ["Famgo", "Nkomo"],
}
tables_cat = ["C1", "C2", "C3", "C4"]
cat_global = {}          # norm(nom) → catégorie (docx1 = référence nationale)
for cat, names in AB.items():
    for n in names:
        cat_global[norm(n)] = cat
for i, tbl in enumerate(cat1):
    for n in tbl["names"]:
        cat_global.setdefault(norm(n), tables_cat[i])

print("\n" + "═" * 70)
print(f"CATÉGORIES (docx national 02062026) : {len(cat_global)} paroisses uniques catégorisées")
matched = {k: v for k, v in cat_global.items() if k in db_index}
print(f"  → correspondance EXACTE avec la BD : {len(matched)} / {len(cat_global)}")
unmatched = [k for k in cat_global if k not in db_index]
print(f"  → sans correspondance BD : {len(unmatched)} — exemples : {unmatched[:12]}")
db_sans_cat = [p["nom"] for k, ps in db_index.items() if k not in cat_global for p in ps]
print(f"  → paroisses BD SANS catégorie : {len(db_sans_cat)} — ex : {db_sans_cat[:8]}")

# Cohérence docx1 (national) vs docx2 (par région)
conflits_12 = 0
c2_index = {}
for e in cat2:
    c2_index.setdefault(norm(e["paroisse"]), set()).add(e["cat"])
for k, cats in c2_index.items():
    g = cat_global.get(k)
    if g and g not in cats:
        conflits_12 += 1
print(f"  → docx2 (par région) : {len(c2_index)} noms uniques ; en CONFLIT de catégorie avec docx1 : {conflits_12}")

# ═══ 3. Croisement FIDÈLES (docx 3) ═══
fid = json.load(open(D + "out_fideles.json", encoding="utf-8"))
fid_uniq = {}
for r in fid:
    fid_uniq.setdefault(norm(r["paroisse"]), r)
m = sum(1 for k in fid_uniq if k in db_index)
print("\n" + "═" * 70)
print(f"FIDÈLES (docx 19_06_2023) : {len(fid)} lignes, {len(fid_uniq)} paroisses uniques")
print(f"  → correspondance exacte BD : {m} / {len(fid_uniq)}")
print(f"  → paroisses du docx ABSENTES de la BD : {len(fid_uniq) - m}")
db_sans_fid = sum(1 for k in db_index if k not in fid_uniq)
print(f"  → paroisses BD absentes du docx : {db_sans_fid}")

# ═══ 4. Croisement NOUVEL EXCEL 04_02_25 ═══
from openpyxl import load_workbook  # noqa: E402
wb = load_workbook(D + "Projet_de_géolocalisation_04_02_25.xlsx", read_only=True, data_only=True)
ws = wb["Paroisses"]
xl_par = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[9]:
        xl_par.append({"nom": str(row[9]).strip(), "district": str(row[8] or ""),
                       "region": str(row[6] or ""), "comm": row[11], "noncomm": row[12]})
ws = wb["Oeuvres"]
xl_oeu = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[10]:
        xl_oeu.append({"nom_oeuvre": str(row[10]).strip(), "paroisse": str(row[2] or "")})
wb.close()

# comparaison avec l'ANCIEN Recap (source BD)
wb = load_workbook(D + "Recap_paroisses_Projet_de_géolocalisation_26mai.xlsx", read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
recap_names = set()
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[5]:
        recap_names.add(norm(str(row[5])))
wb.close()

xl_norm = {norm(p["nom"]) for p in xl_par}
print("\n" + "═" * 70)
print(f"NOUVEL EXCEL 04_02_25 — feuille Paroisses : {len(xl_par)} lignes")
print(f"  → déjà dans l'ancien Recap 26mai (source BD) : {len(xl_norm & recap_names)} / {len(xl_norm)}")
print(f"  → déjà dans la BD : {sum(1 for k in xl_norm if k in db_index)} / {len(xl_norm)}")
print(f"  → NOUVELLES (ni Recap ni BD) : {len([k for k in xl_norm if k not in recap_names and k not in db_index])}")
print(f"Feuille Oeuvres : {len(xl_oeu)} lignes d'œuvres")
oeu_db = {norm(o) for o in Oeuvre.objects.values_list("nom", flat=True)}
oeu_xl = {norm(o["nom_oeuvre"]) for o in xl_oeu}
print(f"  → nom d'œuvre déjà en BD : {len(oeu_xl & oeu_db)} / {len(oeu_xl)}")

json.dump({"db_sans_cat": db_sans_cat, "cat_unmatched": unmatched},
          open(D + "out_match.json", "w", encoding="utf-8"), ensure_ascii=False)
print("\nOK — détails dans out_match.json")
