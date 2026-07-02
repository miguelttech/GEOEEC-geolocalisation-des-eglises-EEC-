# -*- coding: utf-8 -*-
"""Audit phase 4 — correspondance approchée + sémantique des colonnes fidèles."""
import difflib
import json
import os
import re
import sys
import unicodedata

sys.stdout.reconfigure(encoding="utf-8")
import django  # noqa: E402

django.setup()

from apps.geo.models import Paroisse                                          # noqa: E402
from apps.accounts.models import StatistiqueAnnuelle                          # noqa: E402

D = "/app/data_audit/"


def norm(s: str) -> str:
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode()
    s = s.lower()
    s = re.sub(r"\bparoisse( de| du| d')?\b", " ", s)
    s = re.sub(r"\beec\b", " ", s)
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return " ".join(s.split())


db = {norm(p.nom): p for p in Paroisse.objects.select_related("district__region")}
db_keys = list(db.keys())


def fuzzy(k: str, cutoff=0.87):
    if k in db:
        return k, 1.0
    got = difflib.get_close_matches(k, db_keys, n=1, cutoff=cutoff)
    return (got[0], difflib.SequenceMatcher(None, k, got[0]).ratio()) if got else (None, 0)


# ═══ 1. Catégories : exact + fuzzy ═══
cat1 = json.load(open(D + "out_cat1.json", encoding="utf-8"))
match_detail = json.load(open(D + "out_match.json", encoding="utf-8"))
cat_names = set()
for t in cat1:
    cat_names.update(norm(n) for n in t["names"])
AB_names = ["Cinquantenaire", "Messa - Mokolo", "Biyem-Assi", "Nlongkak", "Bonamunduru",
            "Briqueterie 1", "Briqueterie 2", "Essos", "Mont des Oliviers", "New-Bell Aviation",
            "New-Deido", "Njimbam- Foumbot", "Bonabéri", "Bonatéki", "Centenaire", "Company",
            "Melen", "Soboum", "Grand Temple", "Rue Manguiers", "Bafoussam Plateau",
            "Centre Martin Luther", "Maképè Tonnere", "Tamdja", "Bafang II", "Bamendzi A",
            "Bodomo", "Cité – Palmiers", "Koupgouo", "Messassi", "Ngousso Essesselakok",
            "Nkoldongo II", "Nkoldongo I", "Nkozoa", "Nsam Efoulan", "Nzintia", "Oyom Abang",
            "Tchitchap", "Tialong", "Toket", "Famgo", "Nkomo"]
cat_names.update(norm(n) for n in AB_names)

exact = sum(1 for k in cat_names if k in db)
fz = 0
examples = []
for k in cat_names:
    if k in db:
        continue
    g, r = fuzzy(k)
    if g:
        fz += 1
        if len(examples) < 10:
            examples.append(f"{k} → {db[g].nom} ({r:.2f})")
print(f"CATÉGORIES : {len(cat_names)} noms | exact {exact} | +fuzzy {fz} | restants {len(cat_names)-exact-fz}")
print("  Exemples de rapprochements fuzzy :")
for e in examples:
    print(f"    {e}")

# ═══ 2. Fidèles : fuzzy + sémantique n1/n2 ═══
fid = json.load(open(D + "out_fideles.json", encoding="utf-8"))
fid_uniq = {}
for r in fid:
    fid_uniq.setdefault(norm(r["paroisse"]), r)
exact_f, fuzzy_f = 0, 0
paired = []          # (docx_rec, paroisse_bd)
for k, rec in fid_uniq.items():
    if k in db:
        exact_f += 1
        paired.append((rec, db[k]))
    else:
        g, _ = fuzzy(k)
        if g:
            fuzzy_f += 1
            paired.append((rec, db[g]))
print(f"\nFIDÈLES : {len(fid_uniq)} uniques | exact {exact_f} | +fuzzy {fuzzy_f} | introuvables {len(fid_uniq)-exact_f-fuzzy_f}")

# Sémantique : corréler n1/n2 avec communiants/non_communiants de la BD
stats = {s.paroisse_id: s for s in StatistiqueAnnuelle.objects.all()}
agree_n1_comm, agree_n1_noncomm, n_cmp = 0, 0, 0
for rec, p in paired:
    s = stats.get(p.id)
    if not s or rec["n1"] is None or rec["n2"] is None:
        continue
    n_cmp += 1
    # laquelle des 2 colonnes est la plus proche de communiants BD ?
    if abs(rec["n1"] - s.communiants) + abs(rec["n2"] - s.non_communiants) <= \
       abs(rec["n1"] - s.non_communiants) + abs(rec["n2"] - s.communiants):
        agree_n1_comm += 1
    else:
        agree_n1_noncomm += 1
print(f"  Sémantique (sur {n_cmp} paroisses comparées avec stats BD) :")
print(f"    hypothèse n1=communiants / n2=non-communiants : {agree_n1_comm}")
print(f"    hypothèse inverse                              : {agree_n1_noncomm}")

# ═══ 3. Les 2 paroisses « nouvelles » de l'Excel 04_02_25 ═══
from openpyxl import load_workbook  # noqa: E402
wb = load_workbook(D + "Projet_de_géolocalisation_04_02_25.xlsx", read_only=True, data_only=True)
ws = wb["Paroisses"]
news = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[9]:
        k = norm(str(row[9]))
        if k not in db and not fuzzy(k)[0]:
            news.append((row[9], row[8], row[6]))
print(f"\nEXCEL 04_02_25 — paroisses réellement introuvables en BD (même en fuzzy) : {len(set(n[0] for n in news))}")
for n in list(dict.fromkeys(news))[:8]:
    print(f"    {n}")

# ═══ 4. Œuvres du nouvel Excel absentes de la BD (échantillon) ═══
from apps.oeuvres.models import Oeuvre  # noqa: E402
oeu_db_norm = {norm(o) for o in Oeuvre.objects.values_list("nom", flat=True)}
ws = wb["Oeuvres"]
missing_oeu = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[10]:
        k = norm(str(row[10]))
        if k not in oeu_db_norm and not difflib.get_close_matches(k, list(oeu_db_norm), n=1, cutoff=0.85):
            missing_oeu.append(f"{str(row[10]).strip()[:52]}  (paroisse: {row[2]})")
wb.close()
uniq_missing = list(dict.fromkeys(missing_oeu))
print(f"\nŒUVRES du nouvel Excel INTROUVABLES en BD (fuzzy inclus) : {len(uniq_missing)}")
for m in uniq_missing[:15]:
    print(f"    {m}")
json.dump(uniq_missing, open(D + "out_oeuvres_manquantes.json", "w", encoding="utf-8"), ensure_ascii=False)
