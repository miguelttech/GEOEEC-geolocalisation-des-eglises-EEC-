# -*- coding: utf-8 -*-
"""Audit phase 1 — structure de tous les fichiers (nouveaux + anciens)."""
import glob
import sys

sys.stdout.reconfigure(encoding="utf-8")
from docx import Document          # noqa: E402
from openpyxl import load_workbook  # noqa: E402

D = "/app/data_audit/"


def dump_docx(path):
    print(f"\n{'='*80}\nDOCX : {path.split('/')[-1]}\n{'='*80}")
    doc = Document(path)
    # Paragraphes non vides (titres, intro)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    print(f"Paragraphes non vides : {len(paras)} — premiers :")
    for p in paras[:6]:
        print(f"   | {p[:110]}")
    print(f"Tableaux : {len(doc.tables)}")
    for ti, t in enumerate(doc.tables):
        rows = len(t.rows)
        cols = len(t.columns)
        print(f"\n  ── Tableau {ti+1} : {rows} lignes × {cols} colonnes")
        for ri in range(min(4, rows)):          # en-tête + 3 lignes
            cells = [c.text.strip().replace("\n", " ")[:28] for c in t.rows[ri].cells]
            print(f"     L{ri}: {cells}")
        if rows > 4:
            cells = [c.text.strip().replace("\n", " ")[:28] for c in t.rows[rows-1].cells]
            print(f"     L{rows-1} (dernière): {cells}")


def dump_xlsx(path):
    print(f"\n{'='*80}\nXLSX : {path.split('/')[-1]}\n{'='*80}")
    wb = load_workbook(path, read_only=True, data_only=True)
    print(f"Feuilles ({len(wb.sheetnames)}) : {wb.sheetnames}")
    for sn in wb.sheetnames:
        ws = wb[sn]
        print(f"\n  ── Feuille « {sn} » : {ws.max_row} lignes × {ws.max_column} colonnes")
        for ri, row in enumerate(ws.iter_rows(min_row=1, max_row=4, values_only=True)):
            vals = [str(v)[:24] if v is not None else "·" for v in row[:14]]
            print(f"     L{ri+1}: {vals}")
    wb.close()


print("#" * 30, "NOUVEAUX FICHIERS (superviseur)", "#" * 30)
for f in sorted(glob.glob(D + "*.docx")):
    dump_docx(f)
dump_xlsx(D + "Projet_de_géolocalisation_04_02_25.xlsx")

print("\n" + "#" * 30, "ANCIENS FICHIERS (source de la BD)", "#" * 30)
for name in ["Recap_paroisses_Projet_de_géolocalisation_26mai.xlsx",
             "Recap_oeuvres_EEC_2025.xlsx", "OUVRIERS.xlsx"]:
    dump_xlsx(D + name)
