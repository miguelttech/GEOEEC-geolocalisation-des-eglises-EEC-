"""
Inspection complète de toutes les feuilles des fichiers Excel.
But : vérifier qu'on n'a pas manqué de données dans aucun fichier.
"""
import os
import openpyxl

DATA = "/data"
fichiers = os.listdir(DATA)

for f in sorted(fichiers):
    if not f.endswith(".xlsx"):
        continue
    chemin = DATA + "/" + f
    wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    print()
    print("=" * 70)
    print(f"FICHIER : {f}")
    print(f"Feuilles : {wb.sheetnames}")
    print("=" * 70)

    for sh in wb.sheetnames:
        ws = wb[sh]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            print(f"  [{sh}] VIDE")
            continue
        headers = rows[0]
        nb_data = len(rows) - 1
        # Compter les lignes non vides
        non_vides = sum(
            1 for r in rows[1:]
            if any(v is not None and str(v).strip() for v in r)
        )
        print(f"\n  [{sh}] — {nb_data} lignes totales, {non_vides} non vides")
        print(f"    En-têtes : {[str(h)[:25] if h else '' for h in headers[:8]]}")
        # Afficher les 3 premières lignes de données
        for i, row in enumerate(rows[1:4], start=2):
            vals = [str(v)[:20] if v is not None else "" for v in row[:8]]
            print(f"    Ligne {i} : {vals}")

    wb.close()
