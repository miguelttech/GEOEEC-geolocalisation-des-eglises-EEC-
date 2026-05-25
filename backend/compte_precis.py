"""
Comptage precis des vraies entites dans les 3 fichiers Excel.
Distingue : vraies entites / separateurs / lignes vides.
"""
import os
import openpyxl

fichiers = os.listdir("/data")
f_par  = next(f for f in fichiers if "paroisse" in f.lower() or "geo" in f.lower())
f_ouv  = next(f for f in fichiers if "ouvrier" in f.lower())
f_oeuv = next(f for f in fichiers if "oeuvre" in f.lower())


def analyser(filepath, label, col_cle, nom_col_cle):
    """
    Analyse un fichier Excel et distingue :
      - lignes avec une valeur dans la colonne cle (vraies entites)
      - lignes sans valeur dans la colonne cle (separateurs/titres)
      - lignes completement vides
    col_cle : index de la colonne qui identifie une vraie entite (ex: nom paroisse)
    """
    print()
    print("=" * 60)
    print("FICHIER :", label)
    print("=" * 60)

    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)

    for sh_name in wb.sheetnames:
        ws = wb[sh_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        headers = rows[0]
        data = rows[1:]

        print("Feuille :", sh_name)
        print("En-tetes :", [str(h)[:30] if h else "" for h in headers])
        print("Total lignes apres en-tete :", len(data))

        vides = 0
        sans_cle = []
        avec_cle = []

        for i, row in enumerate(data, start=2):
            # Valeurs non nulles dans toute la ligne
            vals = [v for v in row if v is not None and str(v).strip() != ""]

            if not vals:
                vides += 1
            else:
                val_cle = row[col_cle] if len(row) > col_cle else None
                if not val_cle or str(val_cle).strip() == "":
                    sans_cle.append((i, str(vals[:3])))
                else:
                    avec_cle.append(str(val_cle).strip())

        print("  Lignes vides                 :", vides)
        print("  Lignes sans", nom_col_cle, ":", len(sans_cle))
        print("  VRAIES ENTITES               :", len(avec_cle))

        if sans_cle:
            print("  Exemples lignes sans cle :")
            for x in sans_cle[:8]:
                print("    Ligne", x[0], "|", x[1][:80])

    wb.close()


# Paroisse : colonne cle = index 4 (Nom de la paroisse)
analyser("/data/" + f_par, "PAROISSES", 4, "nom paroisse")

# Ouvriers : on va tester plusieurs colonnes pour trouver la cle
print()
print("=" * 60)
print("FICHIER : OUVRIERS (analyse detaillee)")
print("=" * 60)
wb2 = openpyxl.load_workbook("/data/" + f_ouv, read_only=True, data_only=True)
ws2 = wb2.active
rows2 = list(ws2.iter_rows(values_only=True))
headers2 = rows2[0]
data2 = rows2[1:]
print("En-tetes :", [str(h)[:40] if h else "" for h in headers2])
print("Total lignes apres en-tete :", len(data2))
# Trouver la colonne "Nom" (ouvrier)
nom_idx = next((i for i, h in enumerate(headers2) if h and "nom" == str(h).strip().lower()), None)
par_idx = next((i for i, h in enumerate(headers2) if h and "paroisse" in str(h).lower()), None)
print("Colonne 'Nom' (ouvrier) detectee a l index :", nom_idx)
print("Colonne 'Nom de la paroisse' detectee a l index :", par_idx)

vides2, sans_nom2, avec_nom2 = 0, [], []
for i, row in enumerate(data2, start=2):
    vals = [v for v in row if v is not None and str(v).strip() != ""]
    if not vals:
        vides2 += 1
    else:
        # On prend le nom de la paroisse comme cle car chaque ligne = une paroisse
        val = row[par_idx] if par_idx is not None and len(row) > par_idx else None
        if not val or str(val).strip() == "":
            sans_nom2.append((i, str(vals[:3])))
        else:
            avec_nom2.append(str(val).strip())

print("  Lignes vides                         :", vides2)
print("  Lignes sans nom de paroisse          :", len(sans_nom2))
print("  Lignes avec nom de paroisse (entrees):", len(avec_nom2))
if sans_nom2:
    print("  Exemples lignes sans cle :")
    for x in sans_nom2[:5]:
        print("    Ligne", x[0], "|", x[1][:80])
wb2.close()

# Oeuvres : structure differente (une ligne = une region)
print()
print("=" * 60)
print("FICHIER : OEUVRES (analyse detaillee)")
print("=" * 60)
wb3 = openpyxl.load_workbook("/data/" + f_oeuv, read_only=True, data_only=True)
ws3 = wb3.active
rows3 = list(ws3.iter_rows(values_only=True))
headers3 = rows3[0]
data3 = rows3[1:]
print("En-tetes :", [str(h)[:40] if h else "" for h in headers3])
print("Total lignes apres en-tete :", len(data3))

# Colonne Region Synodale = index 0
vides3, sans_region3, avec_region3 = 0, [], []
for i, row in enumerate(data3, start=2):
    vals = [v for v in row if v is not None and str(v).strip() != ""]
    if not vals:
        vides3 += 1
    else:
        region = row[0] if len(row) > 0 else None
        if not region or str(region).strip() == "":
            sans_region3.append((i, str(vals[:4])))
        else:
            avec_region3.append(str(region).strip())

print("  Lignes vides                    :", vides3)
print("  Lignes sans region              :", len(sans_region3))
print("  Lignes avec region (entrees)    :", len(avec_region3))
print("  Regions uniques                 :", len(set(avec_region3)))
print("  Regions :", sorted(set(avec_region3)))
wb3.close()
