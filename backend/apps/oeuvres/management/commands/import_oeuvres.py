"""
=============================================================================
FICHIER N°9 (dans l'ordre de création du projet)
CHEMIN   : backend/apps/oeuvres/management/commands/import_oeuvres.py
CRÉÉ     : Phase 2 — Import des oeuvres de l'EEC
DÉPEND DE: apps/oeuvres/models.py (TypeOeuvre, Oeuvre)
           apps/geo/models.py (RegionSynodale, District, Paroisse)
           import_typeoeuvre.py (les TypeOeuvre doivent exister)
           import_paroisses.py (les paroisses doivent exister)
=============================================================================

OBJECTIF :
  Lire les 3 feuilles du fichier Excel des oeuvres et créer les 311 oeuvres
  de l'EEC dans la base de données, en les associant à leur niveau géographique
  (région, district ou paroisse) et à leur type (scolaire, médical, etc.).

QU'EST-CE QU'UNE "OEUVRE" DANS L'EEC ?
  Une oeuvre est une réalisation concrète de l'Église Évangélique du Cameroun :
  une école, un hôpital, une ferme, un immeuble commercial, un terrain.
  L'EEC distingue 3 niveaux de gestion pour ses oeuvres :
  1. Oeuvres régionales  : gérées par une région synodale (niveau supérieur)
  2. Oeuvres de district  : gérées par un district
  3. Oeuvres paroissiales : gérées par une paroisse individuelle

LE FICHIER SOURCE (3 FEUILLES) :
  Recap_oeuvres_EEC_2025.xlsx
  ┌─────────────────────────┬───────────────────────────────────────────────┐
  │ Nom de la feuille       │ Contenu                                       │
  ├─────────────────────────┼───────────────────────────────────────────────┤
  │ Oeuvres_regionales      │ 59 lignes — oeuvres gérées par une région    │
  │ Oeuvres_districts       │ 55 lignes — oeuvres gérées par un district   │
  │ Oeuvres_paroissialles   │ 158 lignes — oeuvres gérées par une paroisse │
  └─────────────────────────┴───────────────────────────────────────────────┘

CONCEPT CLÉ : LE "PIVOT" (transformation de colonnes en lignes)
  Le fichier Excel n'est PAS structuré "une ligne = une oeuvre".
  Au lieu de ça, chaque ligne peut contenir PLUSIEURS oeuvres de types différents.
  Exemple d'une ligne de la feuille régionale :
    Région    | Scolaire               | PX    | PY    | Médicale       | PX   | PY
    ADAMAOUA  | École Primaire de...   | 7.33  | 13.58 | Hôpital de...  | 7.40 | 13.62

  Une seule ligne → 2 oeuvres (une scolaire, une médicale).
  Pour chaque ligne, on parcourt les "groupes de colonnes" et on crée une oeuvre
  par groupe non vide. Cette transformation s'appelle un "pivot" :
  on transforme des colonnes en lignes de la table Oeuvre.

CORRESPONDANCE GÉOGRAPHIQUE :
  Les noms de régions dans l'Excel peuvent différer légèrement de ceux en base.
  On utilise la même table CORRECTIONS_REGIONS que dans import_paroisses.py.
  Pour les districts et paroisses, on cherche par (region, district) ou (district, paroisse).

COMMANDE D'EXÉCUTION :
  docker compose run --rm backend python manage.py import_oeuvres

OPTIONS :
  --dry-run   : simulation sans écriture en base
  --data-dir  : chemin vers le dossier data (défaut : /data)
=============================================================================
"""

# os : module Python pour les opérations sur les fichiers
import os

# openpyxl : bibliothèque Python pour lire les fichiers Excel
import openpyxl

# BaseCommand : classe de base Django pour les commandes de gestion
from django.core.management.base import BaseCommand

# Point : classe Django GIS pour créer un point géographique (latitude, longitude)
from django.contrib.gis.geos import Point

# Les modèles géographiques (pour retrouver régions, districts, paroisses)
from apps.geo.models import RegionSynodale, District, Paroisse

# Les modèles des oeuvres (qu'on va remplir)
from apps.oeuvres.models import TypeOeuvre, Oeuvre


# =============================================================================
# CONSTANTES DE VALIDATION GÉOGRAPHIQUE
# Même bounding box que dans import_paroisses.py
# =============================================================================
CAM_LAT_MIN, CAM_LAT_MAX = 1.7, 13.1   # Limites nord/sud du Cameroun
CAM_LON_MIN, CAM_LON_MAX = 8.5, 16.2   # Limites est/ouest du Cameroun


# =============================================================================
# TABLE DE CORRECTIONS DES NOMS DE RÉGIONS
# Identique à import_paroisses.py — 5 régions avec des noms différents
# entre l'Excel et la base de données.
# =============================================================================
CORRECTIONS_REGIONS = {
    "HAUT NKAM":           "HAUT-NKAM",
    "HAUTS PLATEAUX":      "HAUTS-PLATEAUX",
    "MOUNGO SUD & MEME":   "MOUNGO SUD ET MEME",
    "NDE MBAM ET INOUBOU": "NDE & MBAM ET INOUBOU",
    "NORD & EXTREME-NORD": "NORD & EXTREME NORD",
}


# =============================================================================
# DICTIONNAIRE DE DÉTECTION DE TYPE PAR MOTS-CLÉS
#
# La colonne "type" de certaines feuilles contient du texte libre.
# Ce dictionnaire permet de détecter le TypeOeuvre depuis ce texte.
# Clé = mot-clé en minuscule → Valeur = code TypeOeuvre
#
# Exemple : si la cellule "type" contient "école primaire"
#   → "école" correspond au mot-clé "école" → TypeOeuvre "SCOLAIRE"
# =============================================================================
MOTS_CLES_TYPE = {
    "scolaire":      "SCOLAIRE",      # le type exact
    "école":         "SCOLAIRE",      # accents courants
    "ecole":         "SCOLAIRE",      # sans accent
    "collège":       "SCOLAIRE",      # collège = établissement scolaire
    "college":       "SCOLAIRE",      # sans accent
    "lycée":         "SCOLAIRE",      # lycée = établissement scolaire
    "lycee":         "SCOLAIRE",      # sans accent
    "universitaire": "UNIVERSITAIRE", # le type exact
    "université":    "UNIVERSITAIRE", # avec accent
    "universite":    "UNIVERSITAIRE", # sans accent
    "médical":       "MEDICALE",      # variante avec accent
    "medical":       "MEDICALE",      # sans accent
    "médicale":      "MEDICALE",      # féminin avec accent
    "medicale":      "MEDICALE",      # féminin sans accent
    "santé":         "MEDICALE",      # centre de santé
    "sante":         "MEDICALE",      # sans accent
    "hôpital":       "MEDICALE",      # hôpital = oeuvre médicale
    "hopital":       "MEDICALE",      # sans accent
    "dispensaire":   "MEDICALE",      # dispensaire = oeuvre médicale
    "centre de":     "MEDICALE",      # "centre de santé"
    "agropastorale": "AGROPASTORALE", # le type exact
    "agropastoral":  "AGROPASTORALE", # variante masculine
    "ferme":         "AGROPASTORALE", # ferme = oeuvre agropastorale
    "immeuble":      "IMMEUBLE",      # le type exact
    "terrain":       "TERRAIN",       # le type exact
}


# =============================================================================
# FONCTIONS UTILITAIRES
# =============================================================================

def corriger_region(nom_excel):
    """
    Normalise le nom d'une région depuis l'Excel vers le nom stocké en base.

    PARAMÈTRE :
      nom_excel : nom brut de la cellule Excel (ex: "HAUT NKAM")

    RETOUR :
      Le nom corrigé (ex: "HAUT-NKAM") ou None si la cellule est vide.

    EXEMPLES :
      corriger_region("HAUT NKAM")    → "HAUT-NKAM"
      corriger_region("ADAMAOUA")     → "ADAMAOUA" (pas de correction nécessaire)
      corriger_region(None)           → None
    """
    if not nom_excel:
        return None

    # str(nom_excel).strip().upper() : nettoyer et mettre en majuscules
    clean = str(nom_excel).strip().upper()

    # CORRECTIONS_REGIONS.get(clean, clean) :
    # Si le nom est dans le dictionnaire → retourner la version corrigée
    # Sinon → retourner le nom tel quel (déjà correct)
    return CORRECTIONS_REGIONS.get(clean, clean)


def point_cameroun(px, py):
    """
    Crée un objet Point GIS si les coordonnées sont dans la bounding box du Cameroun.
    Dans ce fichier, px = latitude et py = longitude (même inversion que import_paroisses.py).

    PARAMÈTRES :
      px : valeur de la colonne Coord_x (= LATITUDE en réalité)
      py : valeur de la colonne Coord_y (= LONGITUDE en réalité)

    RETOUR :
      Un objet Point(longitude, latitude) ou None si coordonnées absentes/invalides.
    """

    # Si l'un des deux est absent → pas de position possible
    if px is None or py is None:
        return None

    try:
        # Convertir en nombres flottants
        # px = Coord_x = LATITUDE (colonne mal nommée dans l'Excel)
        # py = Coord_y = LONGITUDE (colonne mal nommée dans l'Excel)
        lat = float(px)
        lon = float(py)

        # Vérifier que les coordonnées sont dans les limites du Cameroun
        if CAM_LAT_MIN <= lat <= CAM_LAT_MAX and CAM_LON_MIN <= lon <= CAM_LON_MAX:
            # Point(longitude, latitude) : ordre standard GIS (longitude D'ABORD)
            # srid=4326 : système WGS84 (GPS standard)
            return Point(lon, lat, srid=4326)

    except (ValueError, TypeError):
        # Valeur non numérique → pas de position
        pass

    return None  # Coordonnées invalides ou hors Cameroun → pas de position


def detecter_type(texte):
    """
    Détecte le code TypeOeuvre à partir d'un texte libre.
    Cherche les mots-clés définis dans MOTS_CLES_TYPE.

    PARAMÈTRE :
      texte : valeur brute de la cellule "type" dans l'Excel

    RETOUR :
      Le code TypeOeuvre (ex: "SCOLAIRE") ou "AUTRE" si rien ne correspond.
      Retourne None si le texte est vide.

    EXEMPLES :
      detecter_type("Scolaire")          → "SCOLAIRE"
      detecter_type("Hôpital de Bali")   → "MEDICALE"
      detecter_type("Ferme avicole")     → "AGROPASTORALE"
      detecter_type("Inconnu")           → "AUTRE"
      detecter_type(None)                → None
    """

    # Cellule vide → pas de type détecté
    if not texte:
        return None

    # Mettre en minuscules pour une comparaison insensible à la casse
    texte_bas = str(texte).lower()

    # Chercher chaque mot-clé dans le texte
    # Si "école" est dans "École Primaire de Bali" → retourner "SCOLAIRE"
    for mot_cle, code_type in MOTS_CLES_TYPE.items():
        if mot_cle in texte_bas:
            return code_type

    # Aucun mot-clé reconnu → "AUTRE" (catégorie fourre-tout)
    return "AUTRE"


# =============================================================================
# CLASSE PRINCIPALE : Command
# =============================================================================
class Command(BaseCommand):
    """
    Commande Django qui importe les 311 oeuvres depuis les 3 feuilles Excel.
    Django l'exécute avec : python manage.py import_oeuvres
    """

    # Message affiché avec --help
    help = "Importe les oeuvres EEC depuis le fichier Excel (3 feuilles)"

    def add_arguments(self, parser):
        """Déclare les options de ligne de commande."""
        parser.add_argument("--dry-run", action="store_true",
                            help="Simule sans ecrire en base")
        parser.add_argument("--data-dir", default="/data",
                            help="Chemin vers le dossier data (defaut: /data)")

    def handle(self, *args, **options):
        """Méthode principale : ouvre le fichier et traite les 3 feuilles."""

        dry_run  = options["dry_run"]
        data_dir = options["data_dir"]

        if dry_run:
            self.stdout.write(self.style.WARNING("MODE SIMULATION — rien ne sera ecrit"))

        # ═══════════════════════════════════════════════════════════════════
        # ÉTAPE 1 : Trouver le fichier Excel des oeuvres
        # ═══════════════════════════════════════════════════════════════════
        fichiers = os.listdir(data_dir)

        # Chercher le fichier dont le nom contient "oeuvre"
        nom_fichier = next(
            (f for f in fichiers if "oeuvre" in f.lower()), None
        )

        if not nom_fichier:
            self.stderr.write(self.style.ERROR("Fichier oeuvres introuvable dans " + data_dir))
            return

        chemin = os.path.join(data_dir, nom_fichier)
        self.stdout.write(f"Fichier : {nom_fichier}")

        # ═══════════════════════════════════════════════════════════════════
        # ÉTAPE 2 : Charger les caches depuis la base de données
        # ═══════════════════════════════════════════════════════════════════
        # On charge TOUT en mémoire une seule fois pour éviter des milliers
        # de requêtes SQL (une par ligne × par colonne × par feuille).

        # Cache des régions : "NOM_EN_BASE" → objet RegionSynodale
        regions_db = {r.nom.upper(): r for r in RegionSynodale.objects.all()}

        # Cache des districts : ("NOM_REGION", "NOM_DISTRICT") → objet District
        # La clé est un tuple pour gérer les homonymes (même nom dans plusieurs régions)
        districts_db = {
            (d.region.nom.upper(), d.nom.upper()): d
            for d in District.objects.select_related("region").all()
        }

        # Cache des paroisses : ("NOM_DISTRICT", "NOM_PAROISSE") → objet Paroisse
        paroisses_db = {
            (p.district.nom.upper(), p.nom.upper()): p
            for p in Paroisse.objects.select_related("district").all()
        }

        # Cache des types d'oeuvres : "SCOLAIRE" → objet TypeOeuvre
        # (créés par import_typeoeuvre.py)
        types_db = {t.nom: t for t in TypeOeuvre.objects.all()}

        # Vérifier que les TypeOeuvre existent (prérequis)
        if not types_db:
            self.stderr.write(self.style.ERROR(
                "Aucun TypeOeuvre en base ! Lance d'abord : "
                "python manage.py import_typeoeuvre"
            ))
            return

        # Compteurs globaux (somme des 3 feuilles)
        total_crees   = 0
        total_maj     = 0
        total_ignores = 0

        # ═══════════════════════════════════════════════════════════════════
        # ÉTAPE 3 : Ouvrir le fichier Excel
        # ═══════════════════════════════════════════════════════════════════
        # read_only=True : ouverture en lecture seule (performance)
        # data_only=True : récupérer les valeurs calculées, pas les formules
        wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)

        # ═══════════════════════════════════════════════════════════════════
        # ÉTAPE 4 : Traiter les 3 feuilles (une méthode par feuille)
        # ═══════════════════════════════════════════════════════════════════
        # Chaque méthode retourne (nb_crees, nb_maj, nb_ignores)
        # On additionne les résultats dans les compteurs globaux.

        # Feuille 1 : Oeuvres régionales
        self.stdout.write("\n--- Feuille Oeuvres_regionales ---")
        c, m, ig = self._importer_regionales(wb, regions_db, types_db, dry_run)
        total_crees += c; total_maj += m; total_ignores += ig

        # Feuille 2 : Oeuvres de district
        self.stdout.write("\n--- Feuille Oeuvres_districts ---")
        c, m, ig = self._importer_districts(wb, regions_db, districts_db, types_db, dry_run)
        total_crees += c; total_maj += m; total_ignores += ig

        # Feuille 3 : Oeuvres paroissiales
        self.stdout.write("\n--- Feuille Oeuvres_paroissialles ---")
        c, m, ig = self._importer_paroissiales(
            wb, regions_db, districts_db, paroisses_db, types_db, dry_run
        )
        total_crees += c; total_maj += m; total_ignores += ig

        # Fermer le fichier Excel
        wb.close()

        # Bilan global des 3 feuilles
        self.stdout.write("\n" + "=" * 55)
        self.stdout.write(self.style.SUCCESS(
            f"IMPORT OEUVRES TERMINE\n"
            f"  Oeuvres creees      : {total_crees}\n"
            f"  Oeuvres mises a jour: {total_maj}\n"
            f"  Lignes ignorees     : {total_ignores}"
        ))
        self.stdout.write("=" * 55)

    # ===========================================================================
    # MÉTHODE PRIVÉE : _creer_oeuvre()
    # ===========================================================================
    def _creer_oeuvre(self, nom, code_type, position, types_db, dry_run,
                      paroisse=None, district=None, region=None):
        """
        Crée ou met à jour une oeuvre en base de données.

        PARAMÈTRES :
          nom       : nom de l'oeuvre (ex: "Ecole Primaire Catholique de Bonanjo")
          code_type : code du TypeOeuvre (ex: "SCOLAIRE", "MEDICALE")
          position  : objet Point GIS ou None (si pas de GPS)
          types_db  : cache des TypeOeuvre {code: objet}
          dry_run   : True = simuler sans écrire
          paroisse  : objet Paroisse (pour les oeuvres paroissiales)
          district  : objet District (pour les oeuvres de district)
          region    : objet RegionSynodale (pour les oeuvres régionales)

        RETOUR :
          True si créé, False si mis à jour, None si erreur.

        NOTE : Exactement un parmi (paroisse, district, region) doit être fourni.
               Le modèle Oeuvre a ces 3 FK toutes optionnelles (nullable).
        """

        # Tronquer le nom si trop long (max_length=200 dans le modèle Oeuvre)
        # "..." : indiquer visuellement que le nom a été raccourci
        if len(nom) > 197:
            nom = nom[:197] + "..."

        # Récupérer l'objet TypeOeuvre depuis le cache
        # Si le code_type n'est pas trouvé → utiliser "AUTRE" comme fallback
        type_obj = types_db.get(code_type) or types_db.get("AUTRE")

        if not type_obj:
            # Cas très rare : TypeOeuvre "AUTRE" n'existe pas non plus
            return None

        # Mode simulation → afficher sans écrire en base
        if dry_run:
            lieu = paroisse or district or region
            lieu_str = str(lieu) if lieu else "?"
            self.stdout.write(
                f"  [SIM] '{nom[:40]}' | {code_type} | GPS={'oui' if position else 'non'}"
                f" | lieu={lieu_str[:30]}"
            )
            return True  # Compter comme "créé" pour les stats de simulation

        # Créer ou mettre à jour l'oeuvre en base
        try:
            # Construire le dictionnaire de recherche (clé d'unicité)
            # On cherche une oeuvre avec CE nom + CE type + CE lieu géographique
            # Si on trouve → mise à jour ; sinon → création
            lookup = {"nom": nom, "type_oeuvre": type_obj}

            # Ajouter le lien géographique selon le niveau de l'oeuvre
            # Une seule de ces 3 clés est fournie (les autres sont None par défaut)
            if paroisse:
                lookup["paroisse"] = paroisse    # oeuvre paroissiale
            elif district:
                lookup["district"] = district    # oeuvre de district
            elif region:
                lookup["region"] = region        # oeuvre régionale

            # update_or_create : chercher avec le lookup, mettre à jour ou créer
            # defaults={"position"} : seule la position peut changer lors d'une mise à jour
            _, created = Oeuvre.objects.update_or_create(
                **lookup,                        # ** déroule le dict en arguments nommés
                defaults={"position": position}, # mettre à jour la position GPS si elle change
            )
            return created  # True = créé, False = mis à jour

        except Exception as e:
            self.stderr.write(self.style.ERROR(f"  Erreur pour '{nom}' : {e}"))
            return None  # Indiquer une erreur

    # ===========================================================================
    # MÉTHODE PRIVÉE : _importer_regionales()
    # Traite la feuille "Oeuvres_regionales"
    # ===========================================================================
    def _importer_regionales(self, wb, regions_db, types_db, dry_run):
        """
        Importe les oeuvres de la feuille Oeuvres_regionales.

        Structure des colonnes (base 0) :
          [0]  Région
          [1]  PX (latitude du bureau régional — pas utilisé pour les oeuvres)
          [2]  PY (longitude du bureau régional)
          [3]  Type principal (texte libre)
          [4]  Nom de l'oeuvre scolaire
          [5]  PX oeuvre scolaire (latitude)
          [6]  PY oeuvre scolaire (longitude)
          [7]  Nom de l'oeuvre médicale
          [8]  PX oeuvre médicale
          [9]  PY oeuvre médicale
         [10]  Nom de l'immeuble
         [11]  PX immeuble
         [12]  PY immeuble
         [13]  Nom du terrain
         [14]  PX terrain
         [15]  PY terrain
         [16]  Nom de l'oeuvre agropastorale
         [17]  PX agropastorale
         [18]  PY agropastorale
         [19]  Nom de l'oeuvre "Autre"

        Chaque ligne peut contenir jusqu'à 6 oeuvres (une par type).
        """

        # Vérifier que la feuille existe dans le classeur
        if "Oeuvres_regionales" not in wb.sheetnames:
            self.stdout.write(self.style.WARNING("  Feuille Oeuvres_regionales introuvable"))
            return 0, 0, 0  # Retourner 3 zéros (crees, maj, ignores)

        ws   = wb["Oeuvres_regionales"]
        # [1:] : sauter la première ligne (en-têtes)
        rows = list(ws.iter_rows(values_only=True))[1:]

        nb_crees = nb_maj = nb_ig = 0

        # ─── GROUPES DE COLONNES ────────────────────────────────────────────────
        # Chaque "groupe" correspond à un type d'oeuvre et définit quelles colonnes
        # contiennent son nom et ses coordonnées GPS.
        # Format : (code_type, col_nom, col_px, col_py)
        # col_py = None si ce type n'a pas de colonne GPS (comme "AUTRE")
        groupes = [
            ("SCOLAIRE",       4,  5,  6),   # col 4=Nom, 5=PX(lat), 6=PY(lon)
            ("MEDICALE",       7,  8,  9),   # col 7=Nom, 8=PX(lat), 9=PY(lon)
            ("IMMEUBLE",      10, 11, 12),   # col 10=Nom, 11=PX(lat), 12=PY(lon)
            ("TERRAIN",       13, 14, 15),   # col 13=Nom, 14=PX(lat), 15=PY(lon)
            ("AGROPASTORALE", 16, 17, 18),   # col 16=Nom, 17=PX(lat), 18=PY(lon)
            ("AUTRE",         19, None, None), # col 19=Nom, pas de GPS pour "AUTRE"
        ]

        for i, row in enumerate(rows, start=2):
            # Lire le nom de la région (colonne 0)
            nom_region = row[0] if len(row) > 0 else None

            # Ignorer les lignes vides (séparateurs dans l'Excel)
            if not nom_region or str(nom_region).strip() == "":
                nb_ig += 1
                continue

            # Corriger le nom de région et chercher l'objet en base
            nom_region_base = corriger_region(nom_region)
            region_obj = regions_db.get(nom_region_base)

            if not region_obj:
                self.stderr.write(self.style.WARNING(
                    f"  Ligne {i} — Region inconnue : '{nom_region}'"
                ))
                nb_ig += 1
                continue

            # Lire le type principal (colonne 3) — utilisé pour les oeuvres sans nom explicite
            # Exemple : la cellule "type" dit "SCOLAIRE" mais la colonne nom est vide
            # → On génère un nom automatique comme "Scolaire – Région ADAMAOUA"
            type_principal = detecter_type(row[3] if len(row) > 3 else None)

            # ─── PIVOT : parcourir chaque groupe de colonnes ─────────────────────
            # Pour chaque type d'oeuvre possible, vérifier si cette ligne
            # contient une oeuvre de ce type (cellule nom non vide).
            au_moins_une = False  # Pour savoir si on a trouvé au moins une oeuvre dans cette ligne

            for code_type, col_nom, col_px, col_py in groupes:
                # Lire le nom de l'oeuvre pour ce type (si la colonne existe)
                val_nom = row[col_nom] if len(row) > col_nom else None

                # Lire les coordonnées GPS (si les colonnes existent)
                # "col_px and len(row) > col_px" : vérifier que la colonne GPS existe
                px = row[col_px] if col_px is not None and len(row) > col_px else None
                py = row[col_py] if col_py is not None and len(row) > col_py else None

                # Construire le nom de l'oeuvre
                if val_nom and str(val_nom).strip():
                    # Cas 1 : il y a un nom explicite dans la cellule → l'utiliser
                    nom_oeuvre = str(val_nom).strip()
                elif code_type == type_principal:
                    # Cas 2 : pas de nom, mais ce type correspond au type principal de la ligne
                    # → Générer un nom automatique
                    nom_oeuvre = f"{code_type.capitalize()} – Region {region_obj.nom}"
                else:
                    # Cas 3 : pas de nom ET ce n'est pas le type principal
                    # → Cette ligne n'a pas d'oeuvre de ce type → passer au suivant
                    continue

                # Construire le Point GPS (ou None si coordonnées invalides)
                position = point_cameroun(px, py)

                # Créer ou mettre à jour l'oeuvre en base
                created = self._creer_oeuvre(
                    nom_oeuvre, code_type, position, types_db, dry_run,
                    region=region_obj  # lier à la région (pas de paroisse ni district)
                )

                if created is True:
                    nb_crees += 1
                    au_moins_une = True
                elif created is False:
                    nb_maj += 1
                    au_moins_une = True

            # Si aucune oeuvre n'a été trouvée dans cette ligne → compter comme ignorée
            if not au_moins_une:
                nb_ig += 1

        self.stdout.write(f"  Creees={nb_crees} | MaJ={nb_maj} | Ignorees={nb_ig}")
        return nb_crees, nb_maj, nb_ig

    # ===========================================================================
    # MÉTHODE PRIVÉE : _importer_districts()
    # Traite la feuille "Oeuvres_districts"
    # ===========================================================================
    def _importer_districts(self, wb, regions_db, districts_db, types_db, dry_run):
        """
        Importe les oeuvres de la feuille Oeuvres_districts.

        Structure des colonnes (base 0) :
          [0]  Bureau (ignoré)
          [1]  Région
          [2]  District
          [3]  PX bureau (latitude — ignoré pour les oeuvres)
          [4]  PY bureau (longitude — ignoré pour les oeuvres)
          [5]  Type principal (texte libre)
          [6]  Nom oeuvre scolaire        [7]  PX   [8]  PY
          [9]  Nom oeuvre universitaire   [10] PX   [11] PY
         [12]  Nom oeuvre médicale        [13] PX   [14] PY
         [15]  Nom immeuble               [16] PX   [17] PY
         [18]  Nom terrain                [19] PX   [20] PY
         [21]  Nom agropastorale          [22] PX   [23] PY
        """

        if "Oeuvres_districts" not in wb.sheetnames:
            self.stdout.write(self.style.WARNING("  Feuille Oeuvres_districts introuvable"))
            return 0, 0, 0

        ws   = wb["Oeuvres_districts"]
        rows = list(ws.iter_rows(values_only=True))[1:]  # Sauter les en-têtes

        nb_crees = nb_maj = nb_ig = 0

        # Groupes de colonnes pour les oeuvres de district
        groupes = [
            ("SCOLAIRE",       6,  7,  8),
            ("UNIVERSITAIRE",  9, 10, 11),
            ("MEDICALE",      12, 13, 14),
            ("IMMEUBLE",      15, 16, 17),
            ("TERRAIN",       18, 19, 20),
            ("AGROPASTORALE", 21, 22, 23),
        ]

        for i, row in enumerate(rows, start=2):
            # Lire région (col 1) et district (col 2)
            nom_region   = row[1] if len(row) > 1 else None
            nom_district = row[2] if len(row) > 2 else None

            # Ignorer les lignes vides
            if not nom_region or str(nom_region).strip() == "":
                nb_ig += 1
                continue

            # Corriger et chercher la région
            nom_region_base = corriger_region(nom_region)
            region_obj = regions_db.get(nom_region_base)

            if not region_obj:
                nb_ig += 1
                continue

            # Chercher le district via le cache : (NOM_REGION, NOM_DISTRICT)
            # La clé est un tuple pour gérer les districts homonymes entre régions
            cle_dist = (
                nom_region_base,
                str(nom_district).strip().upper() if nom_district else ""
            )
            district_obj = districts_db.get(cle_dist)
            # Si district non trouvé → on rattachera l'oeuvre à la région (fallback)

            # Détecter le type principal (colonne 5)
            type_principal = detecter_type(row[5] if len(row) > 5 else None)

            au_moins_une = False

            for code_type, col_nom, col_px, col_py in groupes:
                val_nom = row[col_nom] if len(row) > col_nom else None
                px      = row[col_px] if len(row) > col_px else None
                py      = row[col_py] if len(row) > col_py else None

                if val_nom and str(val_nom).strip():
                    nom_oeuvre = str(val_nom).strip()
                elif code_type == type_principal:
                    # Générer un nom automatique incluant le district et la région
                    lieu = f"{nom_district} ({nom_region_base})" if nom_district else nom_region_base
                    nom_oeuvre = f"{code_type.capitalize()} – District {lieu}"
                else:
                    continue

                position = point_cameroun(px, py)

                # Stratégie d'association géographique :
                # - Si le district est trouvé → lier au district
                # - Sinon → lier à la région (moins précis mais valide)
                created = self._creer_oeuvre(
                    nom_oeuvre, code_type, position, types_db, dry_run,
                    district=district_obj,
                    region=region_obj if not district_obj else None,
                )

                if created is True:
                    nb_crees += 1; au_moins_une = True
                elif created is False:
                    nb_maj += 1; au_moins_une = True

            if not au_moins_une:
                nb_ig += 1

        self.stdout.write(f"  Creees={nb_crees} | MaJ={nb_maj} | Ignorees={nb_ig}")
        return nb_crees, nb_maj, nb_ig

    # ===========================================================================
    # MÉTHODE PRIVÉE : _importer_paroissiales()
    # Traite la feuille "Oeuvres_paroissialles" (noter la double 'l')
    # ===========================================================================
    def _importer_paroissiales(self, wb, regions_db, districts_db,
                               paroisses_db, types_db, dry_run):
        """
        Importe les oeuvres de la feuille Oeuvres_paroissialles.

        Structure des colonnes (base 0) :
          [0]  Région
          [1]  District
          [2]  Paroisse
          [3]  Type principal (texte libre)
          [4]  Nom oeuvre scolaire        [5]  PX   [6]  PY
          [7]  Nom oeuvre universitaire   [8]  PX   [9]  PY
         [10]  Nom oeuvre médicale        [11] PX   [12] PY
         [13]  Nom immeuble               [14] PX   [15] PY
         [16]  Nom terrain                (pas de GPS pour ce type dans ce fichier)
        """

        # "Oeuvres_paroissialles" avec DEUX 'l' — c'est le nom exact dans le fichier Excel
        if "Oeuvres_paroissialles" not in wb.sheetnames:
            self.stdout.write(self.style.WARNING("  Feuille Oeuvres_paroissialles introuvable"))
            return 0, 0, 0

        ws   = wb["Oeuvres_paroissialles"]
        rows = list(ws.iter_rows(values_only=True))[1:]

        nb_crees = nb_maj = nb_ig = 0
        nb_sans_paroisse = 0  # Oeuvres dont la paroisse n'existe pas en base

        # Groupes de colonnes pour les oeuvres paroissiales
        # Note : TERRAIN n'a pas de GPS dans ce fichier (col_px et col_py = None)
        groupes = [
            ("SCOLAIRE",      4,  5,  6),
            ("UNIVERSITAIRE", 7,  8,  9),
            ("MEDICALE",     10, 11, 12),
            ("IMMEUBLE",     13, 14, 15),
            ("TERRAIN",      16, None, None),  # pas de GPS pour le terrain paroissial
        ]

        for i, row in enumerate(rows, start=2):
            # Lire région (col 0), district (col 1), paroisse (col 2)
            nom_region   = row[0] if len(row) > 0 else None
            nom_district = row[1] if len(row) > 1 else None
            nom_paroisse = row[2] if len(row) > 2 else None

            # Ignorer les lignes sans paroisse
            if not nom_paroisse or str(nom_paroisse).strip() == "":
                nb_ig += 1
                continue

            # Normaliser en majuscules pour la recherche dans le cache
            nom_par_clean  = str(nom_paroisse).strip().upper()
            nom_dist_clean = str(nom_district).strip().upper() if nom_district else ""

            # TENTATIVE 1 : chercher la paroisse par (district + nom)
            paroisse_obj = paroisses_db.get((nom_dist_clean, nom_par_clean))

            if not paroisse_obj:
                # TENTATIVE 2 : chercher uniquement par nom de paroisse
                # Utile si le district est mal orthographié ou absent
                matches = [p for (dist, par), p in paroisses_db.items() if par == nom_par_clean]

                if len(matches) == 1:
                    # Une seule paroisse avec ce nom → c'est forcément la bonne
                    paroisse_obj = matches[0]
                elif len(matches) > 1:
                    # Plusieurs paroisses avec ce nom → prendre la première
                    # et avertir (ambiguïté dans les données)
                    paroisse_obj = matches[0]
                    self.stderr.write(self.style.WARNING(
                        f"  Ligne {i} — Paroisse ambigue '{nom_par_clean}' "
                        f"({len(matches)} resultats) → premiere prise"
                    ))

            # Si toujours pas trouvée → oeuvre sans paroisse, ignorer
            if not paroisse_obj:
                nb_sans_paroisse += 1
                nb_ig += 1
                if not dry_run:
                    self.stderr.write(self.style.WARNING(
                        f"  Ligne {i} — Paroisse non trouvee : '{nom_par_clean}' "
                        f"(district='{nom_dist_clean}')"
                    ))
                continue

            # Détecter le type principal (colonne 3)
            type_principal = detecter_type(row[3] if len(row) > 3 else None)

            au_moins_une = False

            for code_type, col_nom, col_px, col_py in groupes:
                val_nom = row[col_nom] if len(row) > col_nom else None
                # Pour TERRAIN : col_px et col_py sont None → px et py restent None
                px = row[col_px] if col_px is not None and len(row) > col_px else None
                py = row[col_py] if col_py is not None and len(row) > col_py else None

                if val_nom and str(val_nom).strip():
                    # Tronquer à 197 caractères (max_length=200 dans le modèle)
                    nom_oeuvre = str(val_nom).strip()[:197]
                elif code_type == type_principal:
                    # Générer un nom automatique avec le nom de la paroisse
                    nom_oeuvre = f"{code_type.capitalize()} de {paroisse_obj.nom}"
                else:
                    continue

                position = point_cameroun(px, py)

                # Créer l'oeuvre en la liant à sa paroisse
                created = self._creer_oeuvre(
                    nom_oeuvre, code_type, position, types_db, dry_run,
                    paroisse=paroisse_obj  # oeuvre paroissiale : lier à la paroisse
                )

                if created is True:
                    nb_crees += 1; au_moins_une = True
                elif created is False:
                    nb_maj += 1; au_moins_une = True

            if not au_moins_une:
                nb_ig += 1

        # Afficher un avertissement si des paroisses sont introuvables
        if nb_sans_paroisse:
            self.stdout.write(self.style.WARNING(
                f"  {nb_sans_paroisse} lignes ignorees : paroisse non trouvee en base"
            ))

        self.stdout.write(f"  Creees={nb_crees} | MaJ={nb_maj} | Ignorees={nb_ig}")
        return nb_crees, nb_maj, nb_ig
