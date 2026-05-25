"""
Inspection du fichier OUVRIERS.xlsx pour comprendre sa structure.
"""
import os
import openpyxl

DATA = "/data"
f = "OUVRIERS.xlsx"
chemin = DATA + "/" + f

wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
print(f"Feuilles : {wb.sheetnames}")
print()

for sh in wb.sheetnames:
    ws = wb[sh]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print(f"[{sh}] VIDE")
        continue
    headers = rows[0]
    data = rows[1:]
    non_vides = sum(1 for r in data if any(v is not None and str(v).strip() for v in r))
    print("=" * 70)
    print(f"FEUILLE : {sh} — {len(data)} lignes totales, {non_vides} non vides")
    print("=" * 70)

    print("\nEN-TÊTES (tous les index) :")
    for i, h in enumerate(headers):
        nb = sum(1 for r in data if len(r) > i and r[i] is not None and str(r[i]).strip())
        print(f"  [{i:2d}] {str(h)[:50] if h else '(vide)':50s} → {nb} valeurs non vides")

    print("\nPREMIÈRES LIGNES (10 premières) :")
    for i, row in enumerate(data[:10], start=2):
        vals = [str(v)[:20] if v is not None else "" for v in row]
        print(f"  Ligne {i:3d} : {vals}")

    # Valeurs uniques dans la colonne grade (si elle existe)
    # Chercher une colonne grade
    grade_col = None
    for i, h in enumerate(headers):
        if h and "grade" in str(h).lower():
            grade_col = i
            break
    if grade_col is not None:
        grades = {}
        for r in data:
            v = r[grade_col] if len(r) > grade_col else None
            if v:
                k = str(v).strip()
                grades[k] = grades.get(k, 0) + 1
        print(f"\nVALEURS UNIQUES colonne [{grade_col}] (grade) :")
        for g, count in sorted(grades.items(), key=lambda x: -x[1]):
            print(f"  {count:4d}x  '{g}'")

wb.close()
