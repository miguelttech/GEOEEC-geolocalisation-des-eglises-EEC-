"""
Inspection détaillée de la Feuille 1 du fichier paroisses.
Objectif : identifier ce qu'elle contient que Feuil2 n'a pas.
"""
import os
import openpyxl

fichiers = os.listdir("/data")
f_par = next(f for f in fichiers if "paroisse" in f.lower() or "geo" in f.lower())

wb = openpyxl.load_workbook("/data/" + f_par, read_only=True, data_only=True)
ws1 = wb.sheetnames[0]   # première feuille
ws  = wb[ws1]
rows = list(ws.iter_rows(values_only=True))
headers = rows[0]
data = rows[1:]

print("=" * 70)
print(f"FEUILLE : {ws1}")
print(f"Lignes totales : {len(data)}")
print("=" * 70)

print("\nEN-TÊTES COMPLÈTES (tous les index) :")
for i, h in enumerate(headers):
    print(f"  [{i:2d}] {str(h)[:60] if h else '(vide)'}")

# Statistiques sur chaque colonne
print("\nSTATISTIQUES PAR COLONNE (nombre de valeurs non vides) :")
for i, h in enumerate(headers):
    nb = sum(1 for r in data if len(r) > i and r[i] is not None and str(r[i]).strip())
    print(f"  [{i:2d}] {str(h)[:40]:40s} → {nb} valeurs non vides")

# Comparer avec Feuil2
ws2 = wb["Feuil2"]
rows2 = list(ws2.iter_rows(values_only=True))
print()
print("=" * 70)
print("COMPARAISON Feuille1 vs Feuil2")
print("=" * 70)
print(f"  Feuille 1 ('{ws1}') : {len(rows)-1} lignes, {len(headers)} colonnes")
print(f"  Feuil2              : {len(rows2)-1} lignes, {len(rows2[0]) if rows2 else 0} colonnes")

# Paroisses uniques dans Feuille1 (col 5 = Nom de la paroisse)
par_f1 = set()
for r in data:
    nom = r[5] if len(r) > 5 else None
    if nom and str(nom).strip():
        par_f1.add(str(nom).strip().upper())

# Paroisses uniques dans Feuil2 (col 4 = Nom de la paroisse)
par_f2 = set()
for r in rows2[1:]:
    nom = r[4] if len(r) > 4 else None
    if nom and str(nom).strip():
        par_f2.add(str(nom).strip().upper())

print(f"\n  Paroisses uniques Feuille1 : {len(par_f1)}")
print(f"  Paroisses uniques Feuil2   : {len(par_f2)}")
print(f"  Dans Feuille1 mais pas Feuil2 : {len(par_f1 - par_f2)}")
if par_f1 - par_f2:
    print("  Exemples :", list(par_f1 - par_f2)[:10])

# Colonnes statistiques dans Feuille1
print()
print("QUELQUES LIGNES (colonnes 4-15 = données intéressantes) :")
for row in data[:5]:
    vals = [str(v)[:15] if v is not None else "" for v in row[4:16]]
    print(f"  {vals}")

wb.close()
